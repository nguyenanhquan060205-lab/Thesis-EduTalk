"""Vá số liệu mới vào file Word ĐÃ CÓ, giữ nguyên định dạng người dùng đã chỉnh.

KHÔNG sinh lại tài liệu. Chỉ ghi đè phần TEXT của từng ô bảng và từng run trong
đoạn văn — mọi thuộc tính font, màu, khung bảng, hình ảnh giữ nguyên.

    python va_so.py
"""
import json, pathlib, shutil, sys, time, glob
import pandas as pd
from docx import Document

GOC = pathlib.Path(__file__).resolve().parent.parent
F = GOC / "docs" / "BaoCao_SoSanh_research2_research3.docx"

pc = lambda v, n=1: f"{float(v)*100:.{n}f}%"
d_ = lambda v: f"{float(v)*100:+.1f}đ"


def kn(a, b):
    a, b = float(a), float(b)
    return "-" if b >= 1 else f"{(a-b)/(1-b)*100:.1f}%"


def nap(r):
    P = GOC / r / "data" / "processed"
    j = lambda p: json.loads((P / p).read_text(encoding="utf-8"))
    return {"me": j("10_ChotModel/metrics.json"),
            "sth": j("09_TinhChinh/sieu_tham_so.json"),
            "m0": j("04_MocChuan/moc_chuan.json")}


def dat(cell, txt):
    """Ghi text vào ô, GIỮ định dạng của run đầu tiên."""
    txt = str(txt)
    ps = cell.paragraphs
    if not ps:
        return
    p = ps[0]
    if p.runs:
        p.runs[0].text = txt
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)
    for extra in ps[1:]:
        for r in extra.runs:
            r.text = ""


def va_bang(t, hang, tu_dong=1):
    """Ghi các hàng dữ liệu vào bảng, bỏ qua hàng tiêu đề."""
    for i, h in enumerate(hang):
        if tu_dong + i >= len(t.rows):
            break
        cells = t.rows[tu_dong + i].cells
        for j, v in enumerate(h):
            if j < len(cells) and v is not None:
                dat(cells[j], v)


def main():
    if not F.exists():
        sys.exit(f"không thấy {F}")
    bk = F.parent / f".truoc_khi_va_{time.strftime('%H%M%S')}.docx"
    shutil.copy2(F, bk)
    print(f"Sao lưu → {bk.name}\n")

    R2, R3 = nap("research2"), nap("research3")
    d = Document(str(F))
    T = d.tables
    n = 0

    # ── #3 — bảng tóm tắt 1.4 ────────────────────────────────────────────
    h = []
    for k in ("1", "2", "3", "5"):
        o = [f"Top-{k}"]
        for X in (R2, R3):
            t = X["me"]["test"]["tu_van"]
            a, b = t["top"][k], t["bua"][k]
            o += [pc(a), pc(b), d_(a - b), kn(a, b)]
        h.append(o)
    va_bang(T[3], h); n += 1

    # ── #7 — mốc chuẩn M₀ ────────────────────────────────────────────────
    va_bang(T[7], [
        ["M₀ tư vấn Top-3", pc(R2["m0"]["guided"]["top"]["3"]),
         pc(R3["m0"]["guided"]["top"]["3"])],
        ["đoán bừa", pc(R2["m0"]["guided"]["bua"]["3"]),
         pc(R3["m0"]["guided"]["bua"]["3"])],
        ["M₀ khám phá Top-3", pc(R2["m0"]["auto"]["top"]["3"]),
         pc(R3["m0"]["auto"]["top"]["3"])],
        ["Nhóm đạt 80%",
         f"{sum(1 for x in R2['m0']['theo_khoi'] if x['top3']>=.8)}/{len(R2['m0']['theo_khoi'])}",
         f"{sum(1 for x in R3['m0']['theo_khoi'] if x['top3']>=.8)}/{len(R3['m0']['theo_khoi'])}"],
    ]); n += 1

    # ── #12 — kết quả NB09 ───────────────────────────────────────────────
    va_bang(T[12], [
        ["Top-3 sau tinh chỉnh", pc(R2["sth"]["ket_qua_cv"]["guided_top3"]),
         pc(R3["sth"]["ket_qua_cv"]["guided_top3"])],
        ["Chênh train − val", d_(R2["sth"]["ket_qua_cv"]["gap_top3"]),
         d_(R3["sth"]["ket_qua_cv"]["gap_top3"])],
        ["Trọng số dòng thật", f"×{R2['sth']['w_that']:.0f}",
         f"×{R3['sth']['w_that']:.0f}"],
        ["Lợi ích phân tầng", d_(R2["sth"]["loi_ich_phan_tang"]),
         d_(R3["sth"]["loi_ich_phan_tang"])],
    ]); n += 1

    # ── #13 — kết quả NB10 ───────────────────────────────────────────────
    va_bang(T[13], [
        ["TEST tư vấn Top-3", pc(R2["me"]["test"]["tu_van"]["top"]["3"]),
         pc(R3["me"]["test"]["tu_van"]["top"]["3"])],
        ["đoán bừa", pc(R2["me"]["test"]["tu_van"]["bua"]["3"]),
         pc(R3["me"]["test"]["tu_van"]["bua"]["3"])],
        ["TEST khám phá Top-3", pc(R2["me"]["test"]["kham_pha"]["top"]["3"]),
         pc(R3["me"]["test"]["kham_pha"]["top"]["3"])],
        ["macro-F1", f"{R2['me']['test']['auto_macro_f1']:.3f}",
         f"{R3['me']['test']['auto_macro_f1']:.3f}"],
    ]); n += 1

    # ── #14,15 (research2) và #18,19 (research3) — TRAIN/TEST ────────────
    for idx, X, kh in ((14, R2, "tu_van"), (15, R2, "kham_pha"),
                       (18, R3, "tu_van"), (19, R3, "kham_pha")):
        t_ = X["me"]["test"][kh]; tr = X["me"]["tren_train"][kh]
        va_bang(T[idx], [[f"Top-{k}", pc(tr[k]), pc(t_["top"][k]),
                          d_(tr[k] - t_["top"][k]), pc(t_["bua"][k]),
                          d_(t_["top"][k] - t_["bua"][k]),
                          kn(t_["top"][k], t_["bua"][k])]
                         for k in ("1", "2", "3", "5")]); n += 1

    # ── #16, #20 — bảng đối chứng ────────────────────────────────────────
    for idx, X in ((16, R2), (20, R3)):
        te = X["me"]["test"]; mc = te["kham_pha"]["top"]["3"]
        va_bang(T[idx], [
            ["Đoán lớp đông nhất", pc(te["dong_nhat"]["3"]), d_(te["dong_nhat"]["3"] - mc)],
            ["Model phẳng (không dùng nhóm)", pc(te["phang"]["3"]), d_(te["phang"]["3"] - mc)],
            ["Chỉ 574 phiếu khảo sát", pc(te["chi_khao_sat"]["kham_pha"]["3"]),
             d_(te["chi_khao_sat"]["kham_pha"]["3"] - mc)],
            ["★ Mô hình cuối", pc(mc), "—"],
        ]); n += 1

    # ── #17, #21 — bóc tách theo nhóm ────────────────────────────────────
    for idx, X in ((17, R2), (21, R3)):
        tk = pd.DataFrame(X["me"]["theo_khoi"]).sort_values("n_nganh", ascending=False)
        va_bang(T[idx], [[r.khoi, r.n_nganh, r.n_sv, pc(r.top3), pc(r.bua3),
                          d_(r.hon), kn(r.top3, r.bua3)] for r in tk.itertuples()])
        n += 1

    # ── Đoạn văn có số rời ───────────────────────────────────────────────
    tv2 = R2["me"]["test"]["tu_van"]; tv3 = R3["me"]["test"]["tu_van"]
    kp2 = R2["me"]["test"]["kham_pha"]; kp3 = R3["me"]["test"]["kham_pha"]
    THAY = [
        ("research3 đạt Top-3 = 81.4%", f"research3 đạt Top-3 = {pc(tv3['top']['3'])}"),
        ("research2 đạt 65.7%", f"research2 đạt {pc(tv2['top']['3'])}"),
    ]
    n_p = 0
    for p in d.paragraphs:
        for r in p.runs:
            for cu, moi in THAY:
                if cu in r.text:
                    r.text = r.text.replace(cu, moi); n_p += 1

    d.save(str(F))
    print(f"✅ Đã vá {n} bảng · {n_p} đoạn văn")
    print(f"   Giữ nguyên: định dạng, màu, khung bảng, {len(d.inline_shapes)} hình")
    print(f"\n   research2  test Top-3 = {pc(tv2['top']['3'])}"
          f"   ·   train = {pc(R2['me']['tren_train']['tu_van']['3'])}")
    print(f"   research3  test Top-3 = {pc(tv3['top']['3'])}"
          f"   ·   train = {pc(R3['me']['tren_train']['tu_van']['3'])}")


if __name__ == "__main__":
    main()
