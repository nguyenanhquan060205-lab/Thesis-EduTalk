"""Phần dùng chung cho ba báo cáo Word: Hướng 1 (research/), Hướng 2 (research1/), so sánh.

Ba nguyên tắc:
  1. Không gõ tay con số nào — mọi số đọc từ file kết quả của pipeline hoặc file dữ liệu thô.
     Số hiển thị lấy từ file CSV đủ độ chính xác; file JSON đã làm tròn 4 chữ số, làm tròn
     thêm lần nữa sẽ lệch (0,75648 lưu thành 0,7565 rồi hiện 75,7% thay vì 75,6%).
  2. Số viết kiểu Việt Nam (90,6% · 0,583 · 16.296) và MỌI số in ra đều được ghi sổ.
     Sau khi sinh Word, `kiem_so()` quét tài liệu: số nào không có trong sổ là số gõ tay.
  3. Times New Roman 13 cho mọi chữ, kể cả tiêu đề và bảng — đúng định dạng người dùng đã
     chỉnh trên bản Word.
"""
import hashlib
import json
import pathlib
import re
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GOC = pathlib.Path(__file__).resolve().parent.parent
FONT, CO = "Times New Roman", 13
DEN = RGBColor(0, 0, 0)
TOP_K = (1, 2, 3, 5)

# ═══════════════════════════════════════════════════════════════════════════
#  SỐ — định dạng kiểu Việt Nam và ghi sổ
# ═══════════════════════════════════════════════════════════════════════════
SO_DA_IN = set()


def _q(x, n, nhan=1):
    d = Decimal(repr(round(float(x), 10))) * nhan
    return d.quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP)


def _viet(d):
    s = f"{abs(d):,}".replace(",", "_").replace(".", ",").replace("_", ".")
    SO_DA_IN.add(s)
    return s


def so(x, n=1):
    """0.5829 → '0,583' (n=3)."""
    d = _q(x, n)
    return ("−" if d < 0 else "") + _viet(d)


def nguyen(x):
    """16296 → '16.296'."""
    return so(x, 0)


def pt(x, n=1):
    """Tỉ lệ 0–1 → '90,6%'."""
    d = _q(x, n, 100)
    return ("−" if d < 0 else "") + _viet(d) + "%"


def diem(x, n=1, dau=True):
    """Hiệu hai tỉ lệ → điểm phần trăm.

    dau=True  → '+4,4' / '−2,1', dùng trong cột "chênh" của bảng.
    dau=False → '4,4', dùng trong câu đã có chữ "cao hơn / thấp hơn".
    """
    d = _q(x, n, 100)
    if not dau:
        assert d >= 0, "số điểm trong câu văn phải dương — đảo chiều câu"
        return _viet(d)
    return ("+" if d >= 0 else "−") + _viet(d)


def sai_so_95(p, n):
    """Nửa khoảng tin cậy 95% của một tỉ lệ (dạng tỉ lệ 0–1)."""
    return 1.96 * (p * (1 - p) / n) ** 0.5


# Số được phép xuất hiện mà không qua sổ: tên gọi, không phải số liệu
BO_QUA = [r"[Gg]iai đoạn \d+(?:\s*[–-]\s*\d+)?", r"Hình \d+", r"[Mm]ục \d+(?:\.\d+)*",
          r"Top-\d+", r"Hit@\d+", r"NDCG@\d+", r"@\d+", r"SHA-\d+", r"MD5", r"CRC32",
          r"FIPS 180-4", r"\b(?:19|20)\d\d\b", r"\b[A-Z]\d\d\b", r"XGBoost", r"F1",
          r"ga_\d+", r"Hướng [12]", r"\bH[12]\b"]


def kiem_so(duong):
    """Trả về các số có trong tài liệu nhưng không đi qua sổ — tức số gõ tay."""
    d = Document(str(duong))
    doan = [(p.text, p.style.name) for p in d.paragraphs]
    doan += [(c.text, "bảng") for t in d.tables for r in t.rows for c in r.cells]
    la = []
    for txt, kieu in doan:
        if kieu.startswith("Heading"):
            continue
        t = txt
        for m in BO_QUA:
            t = re.sub(m, " ", t)
        for tok in re.findall(r"\d[\d.,]*\d|\d", t):
            if tok not in SO_DA_IN:
                la.append((tok, txt[:100]))
    return la


# ═══════════════════════════════════════════════════════════════════════════
#  KẾT QUẢ CỦA MỘT PIPELINE
# ═══════════════════════════════════════════════════════════════════════════
class KetQua:
    """Đọc thẳng từ <kho>/data/processed/ và <kho>/data/raw/."""

    def __init__(self, kho):
        self.kho = kho
        self.P = GOC / kho / "data" / "processed"
        self.RAW = GOC / kho / "data" / "raw"
        self.tm = {int(p.name[:2]): p for p in self.P.iterdir()
                   if p.is_dir() and p.name[:2].isdigit()}

    def f(self, gd, ten):
        p = self.tm[gd] / ten
        assert p.exists(), f"thiếu {p}"
        return p

    def js(self, gd, ten):
        return json.loads(self.f(gd, ten).read_text(encoding="utf-8"))

    def cv(self, gd, ten, **kw):
        return pd.read_csv(self.f(gd, ten), **kw)

    def hinh(self, gd, mau):
        g = sorted(self.tm[gd].glob(mau))
        assert g, f"thiếu hình {mau} ở giai đoạn {gd}"
        return g[0]

    def bam(self, gd, ten):
        return hashlib.sha256(self.f(gd, ten).read_bytes()).hexdigest()


def nen(kq):
    """Số liệu giai đoạn 1–7 — dùng chung cho mọi báo cáo, có kiểm chéo."""
    m = kq.js(1, "mapping.json")
    ier = kq.cv(1, "bang_ier.csv")
    ks = kq.cv(1, "khaosat_sach.csv")
    tt = kq.cv(2, "ttth_sach.csv", usecols=["ma_nganh"])
    t3, t4 = kq.js(3, "tom_tat.json"), kq.js(4, "tom_tat.json")
    tho_ks = pd.read_csv(kq.RAW / "khao_sat_dinh_huong_nganh_hoc_raw.csv")
    tho_tt = pd.read_excel(kq.RAW / "TTTH_THPT_raw.xlsx", usecols=["KQ"])
    kd = kq.cv(5, "bao_cao_kiem_dinh.csv").set_index("ban")
    kl = kq.js(5, "ket_luan.json")
    tap = {x: kq.cv(7, f"{x}.csv", usecols=["ma_nganh", "ma_nhom", "nguon"])
           for x in ("train", "val", "test_KHOA")}

    n_likert = int(t3["n_likert"])
    loai = ier[ier.bi_loai == 1]
    # Mọi phiếu bị loại đều chọn cùng một mức cho cả 10 câu — nên được phép mô tả gọn
    # là "chọn cùng một mức cho cả 10 câu" thay vì kể ba chỉ số IER.
    assert len(ier) == len(tho_ks) and (loai.longstring >= n_likert).all()
    assert len(ks) == len(tho_ks) - len(loai) == int(m["n_sach"])
    nganh_nhom = {int(a): int(b) for a, b in m["nganh_to_nhom"].items()}
    co_moi_nhom = Counter(nganh_nhom.values())
    co_nhom = sorted(co_moi_nhom.values(), reverse=True)
    assert co_nhom == sorted((len(v) for v in m["nhom_chi_tiet"].values()), reverse=True)
    vc = ks.ma_nganh.value_counts()
    n_nganh = len(m["ma_to_ten"])
    n_nganh_tt = int(tt.ma_nganh.nunique())
    assert int(t4["n_tu_ttth"]) == len(tt) and int(t4["n_nganh"]) == n_nganh
    n_thieu = n_nganh - n_nganh_tt
    assert int(t4["n_bootstrap"]) % n_thieu == 0
    assert kl["phuong_phap_chot"] in kd.index and bool(kd.loc[kl["phuong_phap_chot"], "hop_le"])
    lk = ks[m["cot_likert"]]
    assert len(m["cot_likert"]) == n_likert
    return dict(
        mapping=m, ten=m["ma_to_ten"], ten_nhom={int(k): v for k, v in m["ten_nhom"].items()},
        nganh_nhom=nganh_nhom, co_moi_nhom=co_moi_nhom, co_nhom=co_nhom,
        n_ks_tho=len(tho_ks), n_ks_loai=len(loai), n_ks=len(ks), n_likert=n_likert,
        likert_min=int(lk.min().min()), likert_max=int(lk.max().max()),
        n_nganh=n_nganh, n_nhom=len(m["ten_nhom"]),
        ks_it_nhat=int(vc.min()), ks_it_nhat_ten=m["ma_to_ten"][str(vc.idxmin())],
        n_tt_tho=len(tho_tt), n_tt=len(tt), n_nganh_tt=n_nganh_tt, n_thieu=n_thieu,
        K=int(t3["K_co_ngot"]), n_bootstrap=int(t4["n_bootstrap"]),
        bs_moi_nganh=int(t4["n_bootstrap"]) // n_thieu, the_he=t4["ds_the_he_ga"],
        kd=kd, kl=kl, tap=tap,
        nguon=lambda x, ng: int((tap[x].nguon == ng).sum()),
    )


def bang_kq(kq):
    """{tên tập: {k: tỉ lệ}} đủ độ chính xác, từ bang_ket_qua.csv của Giai đoạn 10."""
    b = kq.cv(10, "bang_ket_qua.csv").set_index("tap")
    return {t: {int(c[3:]): float(b.loc[t, c]) for c in b.columns} for t in b.index}


def bang_moc(kq):
    """Bốn mốc + mô hình đối chiếu, đủ độ chính xác, từ bang_moc.csv của Giai đoạn 8."""
    b = kq.cv(8, "bang_moc.csv")
    return [r for r in b.to_dict("records")]


def bua_trong_nhom(N, ma_nganh, k):
    """Bốc ngẫu nhiên k ngành trong nhóm đúng."""
    nh, co = N["nganh_nhom"], N["co_moi_nhom"]
    return float(np.mean([min(k, co[nh[int(x)]]) / co[nh[int(x)]] for x in ma_nganh]))


def pho_bien(hoc, do, k):
    """Quy tắc 'luôn gợi ý k ngành đông thí sinh nhất trong nhóm': học tần suất từ `hoc`,
    đo trên `do`."""
    top = {g: set(s.value_counts().head(k).index) for g, s in hoc.groupby("ma_nhom").ma_nganh}
    return sum(n in top[g] for n, g in zip(do.ma_nganh, do.ma_nhom)) / len(do)


def moc_bua(N, hoc, do, M_bua):
    """Ba mốc đoán bừa đủ độ chính xác, đối chiếu với bản làm tròn trong metrics.json."""
    n = N["n_nganh"]
    ra = {"trong_nhom": {k: bua_trong_nhom(N, do.ma_nganh, k) for k in TOP_K},
          "dong_nhat": {k: pho_bien(hoc, do, k) for k in TOP_K},
          "kham_pha": {k: min(k, n) / n for k in TOP_K}}
    for ten, khoa in [("trong_nhom", "tu_van_trong_nhom"), ("dong_nhat", "tu_van_nganh_dong_nhat"),
                      ("kham_pha", "kham_pha")]:
        for k in TOP_K:
            assert abs(ra[ten][k] - M_bua[khoa][f"top{k}"]) < 6e-5, (ten, k)
    return ra


class MoHinhDaLuu:
    """Mô hình cuối đã lưu ở Giai đoạn 10, dùng để tính Top-k mà pipeline không lưu (Top-4).

    Tái hiện đúng MoHinhNganh.diem(): tầng khám phá dùng mô hình chung, tầng tư vấn dùng mô
    hình riêng của nhóm rồi che các ngành ngoài nhóm. Script gọi phải đối chiếu kết quả với
    bang_ket_qua.csv ở các k pipeline có lưu, trước khi tin số Top-4.
    """

    def __init__(self, kq):
        from xgboost import XGBClassifier

        d = kq.tm[10]
        info = json.loads((d / "lop_va_dac_trung.json").read_text(encoding="utf-8"))
        self.cot = info["cot_dac_trung"]
        self.nganh = np.array(info["nganh"])
        assert list(self.nganh) == sorted(self.nganh) and info["rieng_nhom"] is True
        self.nhom = {int(k): int(v) for k, v in info["nganh_to_nhom"].items()}
        self.vi_tri = {int(m): j for j, m in enumerate(self.nganh)}
        self.chung = XGBClassifier()
        self.chung.load_model(d / "model_nganh.json")
        self.rieng = {}
        for g in sorted(set(self.nhom.values())):
            m = XGBClassifier()
            m.load_model(d / f"model_nhom{g}.json")
            self.rieng[g] = (m, sorted(n for n in self.nganh if self.nhom[int(n)] == g))

    def diem(self, df, tu_van):
        X = df[self.cot].values
        S = self.chung.predict_proba(X)
        if not tu_van:
            return S
        nhom = df.ma_nhom.values
        S = S.copy()
        for g, (m, ds) in self.rieng.items():
            sel = nhom == g
            if sel.any():
                pr = m.predict_proba(X[sel])
                T = np.zeros((int(sel.sum()), len(self.nganh)))
                for c, ma in enumerate(ds):
                    T[:, self.vi_tri[int(ma)]] = pr[:, c]
                S[sel] = T
        che = np.array([[self.nhom[int(n)] == g for n in self.nganh] for g in nhom], float)
        return S * che

    def top_k(self, df, tu_van, ks):
        o = np.argsort(-self.diem(df, tu_van), axis=1)
        y = df.ma_nganh.values
        return {k: float(np.mean([y[i] in self.nganh[o[i, :k]] for i in range(len(y))]))
                for k in ks}


def cau_hinh_chot(QUET, hp):
    """Dòng của cấu hình đã chốt trong ket_qua_quet.csv (đủ độ chính xác)."""
    chon = np.ones(len(QUET), dtype=bool)
    for k, v in hp.items():
        chon &= np.isclose(QUET[k].astype(float), float(v))
    assert chon.sum() == 1
    return QUET[chon].iloc[0]


# ═══════════════════════════════════════════════════════════════════════════
#  WORD — Times New Roman 13 cho mọi chữ
# ═══════════════════════════════════════════════════════════════════════════
class Word:
    def __init__(self):
        self.d = Document()
        for ten in ("Normal", "List Bullet", "List Number", "Heading 1", "Heading 2",
                    "Heading 3"):
            st = self.d.styles[ten]
            st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(CO), DEN
            rpr = st.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:          # bỏ font theo theme, nếu không Word tự đổi font
                for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    fonts.attrib.pop(qn(a), None)
                fonts.set(qn("w:eastAsia"), FONT)
        pf = self.d.styles["Normal"].paragraph_format
        pf.space_after, pf.line_spacing = Pt(6), 1.25
        for s in self.d.sections:
            s.left_margin = s.right_margin = Cm(2.2)
            s.top_margin = s.bottom_margin = Cm(2.0)
        self.n_hinh = 0
        self.n_bang = 0

    def _run(self, par, txt, dam=False, ngh=False):
        r = par.add_run(txt)
        r.bold, r.italic = dam, ngh
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(CO), DEN
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        return r

    def h(self, txt, lv):
        p = self.d.add_heading("", level=lv)
        self._run(p, txt, dam=True)
        return p

    def p(self, *doan, dam=False, ngh=False, giua=False):
        """`doan` là chuỗi, hoặc cặp (chuỗi, in_đậm) để in đậm một phần câu."""
        pr = self.d.add_paragraph()
        for x in doan:
            t, d = (x, dam) if isinstance(x, str) else x
            self._run(pr, t, dam=d, ngh=ngh)
        if giua:
            pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return pr

    def gach(self, muc, kieu="List Bullet"):
        """Mỗi mục là chuỗi, hoặc cặp (phần in đậm, phần còn lại)."""
        for x in muc:
            pr = self.d.add_paragraph(style=kieu)
            if isinstance(x, str):
                self._run(pr, x)
            else:
                self._run(pr, x[0], dam=True)
                self._run(pr, x[1])

    def so_thu_tu(self, muc):
        self.gach(muc, kieu="List Number")

    def bang(self, dau, hang, rong=None):
        self.n_bang += 1
        t = self.d.add_table(rows=1, cols=len(dau))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(dau):
            self._run(t.rows[0].cells[i].paragraphs[0], str(c), dam=True)
        for h_ in hang:
            cells = t.add_row().cells
            for i, v in enumerate(h_):
                self._run(cells[i].paragraphs[0], "" if v is None else str(v), dam=(i == 0))
        if rong:
            for i, w in enumerate(rong):
                for row in t.rows:
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph()
        return t

    def hinh(self, duong, chu, rong=16.0):
        self.n_hinh += 1
        self.d.add_picture(str(duong), width=Cm(rong))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.p(f"Hình {self.n_hinh}. {chu}", ngh=True, giua=True)

    def ngat(self):
        self.d.add_page_break()

    def luu(self, duong):
        duong.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(str(duong))
        la = kiem_so(duong)
        print(f"✅ {duong.name} · {self.n_hinh} hình · {self.n_bang} bảng")
        if la:
            print(f"   ⚠️  {len(la)} số không qua sổ (có thể là gõ tay):")
            for tok, ngu_canh in la:
                print(f"      {tok!r:>10}  ←  {ngu_canh}")
        else:
            print("   ✓ mọi con số trong tài liệu đều đọc từ file kết quả")
        return la
